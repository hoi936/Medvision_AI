"""Portable report exports with the exact visual evidence sent to Hermes."""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Any, Iterable

import numpy as np
from PIL import Image


IMAGE_LABELS = {
    "original": "Ảnh X-quang gốc đã preprocess",
    "heatmap": "Grad-CAM heatmap",
    "overlay": "Ảnh X-quang chồng Grad-CAM",
    "pseudo_bbox": "Pseudo bbox suy ra từ Grad-CAM",
}


@dataclass(frozen=True)
class ReportExport:
    data: bytes
    filename: str
    mime: str


def image_to_png_bytes(image: Any) -> bytes:
    """Encode an ndarray/PIL image as a normalized, portable PNG."""
    if isinstance(image, Image.Image):
        pil_image = image.copy()
    else:
        array = np.asarray(image)
        if array.dtype != np.uint8:
            if np.issubdtype(array.dtype, np.floating) and array.size and array.max() <= 1:
                array = array * 255
            array = np.clip(array, 0, 255).astype(np.uint8)
        pil_image = Image.fromarray(array)
    if pil_image.mode not in ("L", "RGB", "RGBA"):
        pil_image = pil_image.convert("RGB")
    buffer = io.BytesIO()
    pil_image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def build_visual_evidence(
    *, original: Any, findings: Iterable[tuple[dict[str, Any], dict[str, Any]]]
) -> list[dict[str, Any]]:
    """Freeze the images used for one report so later UI changes cannot alter it."""
    original_png = image_to_png_bytes(original)
    evidence: list[dict[str, Any]] = []
    for finding, visual in findings:
        evidence.append(
            {
                "finding": str(finding["finding"]),
                "class_id": int(finding["class_id"]),
                "score": float(finding["score"]),
                "threshold": float(finding["threshold"]),
                "decision": "POSITIVE" if finding["positive"] else "NEGATIVE",
                "images": {
                    "original": original_png,
                    "heatmap": image_to_png_bytes(visual["heatmap"]),
                    "overlay": image_to_png_bytes(visual["overlay"]),
                    "pseudo_bbox": image_to_png_bytes(visual["boxed_image"]),
                },
            }
        )
    if not evidence:
        evidence.append(
            {
                "finding": "Không có finding vượt threshold",
                "class_id": None,
                "score": None,
                "threshold": None,
                "decision": "NO_POSITIVE_FINDING",
                "images": {"original": original_png},
            }
        )
    return evidence


def visual_evidence_digest(evidence: Iterable[dict[str, Any]]) -> str:
    import hashlib

    digest = hashlib.sha256()
    for item in evidence:
        digest.update(str(item.get("finding", "")).encode("utf-8"))
        for kind, png in item.get("images", {}).items():
            digest.update(kind.encode("ascii"))
            digest.update(png)
    return digest.hexdigest()


def _visual_markdown(evidence: Iterable[dict[str, Any]]) -> str:
    blocks = [
        "## Phụ lục hình ảnh",
        "",
        "> Grad-CAM và pseudo bbox chỉ thể hiện vùng mô hình chú ý; không phải "
        "segmentation, annotation hoặc xác nhận tổn thương của bác sĩ.",
    ]
    for item in evidence:
        blocks.extend(["", f"### {item['finding']}", ""])
        for kind, png in item["images"].items():
            label = IMAGE_LABELS[kind]
            data = base64.b64encode(png).decode("ascii")
            blocks.extend([f"**{label}**", "", f"![{label}](data:image/png;base64,{data})", ""])
    return "\n".join(blocks).rstrip()


def _split_review_gate(report: str) -> tuple[str, str]:
    required = (
        "review_status: PENDING_CLINICIAN_REVIEW",
        "requires_doctor_review: true",
    )
    lines = report.rstrip().splitlines()
    if len(lines) >= 2 and tuple(line.strip() for line in lines[-2:]) == required:
        return "\n".join(lines[:-2]).rstrip(), "\n".join(required)
    return report.rstrip(), ""


def export_markdown(report: str, evidence: Iterable[dict[str, Any]], filename: str) -> ReportExport:
    body, review_gate = _split_review_gate(report)
    content = body + "\n\n" + _visual_markdown(evidence)
    if review_gate:
        content += "\n\n" + review_gate
    content += "\n"
    return ReportExport(content.encode("utf-8"), f"{filename}.md", "text/markdown")


def _append_docx_markdown(document: Any, report: str) -> None:
    for raw_line in report.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line.replace("**", "").replace("`", ""))


def export_docx(report: str, evidence: Iterable[dict[str, Any]], filename: str) -> ReportExport:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError("Thiếu python-docx; hãy cài lại requirements.txt.") from exc

    document = Document()
    report_body, review_gate = _split_review_gate(report)
    _append_docx_markdown(document, report_body)
    document.add_page_break()
    document.add_heading("Phụ lục hình ảnh", level=1)
    document.add_paragraph(
        "Grad-CAM và pseudo bbox chỉ thể hiện vùng mô hình chú ý; không phải "
        "segmentation, annotation hoặc xác nhận tổn thương của bác sĩ."
    )
    for item in evidence:
        document.add_heading(item["finding"], level=2)
        for kind, png in item["images"].items():
            document.add_paragraph(IMAGE_LABELS[kind])
            document.add_picture(io.BytesIO(png), width=Inches(5.8))
    if review_gate:
        document.add_heading("Trạng thái duyệt", level=2)
        _append_docx_markdown(document, review_gate)
    output = io.BytesIO()
    document.save(output)
    return ReportExport(
        output.getvalue(),
        f"{filename}.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _register_pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    for path in candidates:
        if path.is_file():
            pdfmetrics.registerFont(TTFont("MedVisionUnicode", str(path)))
            return "MedVisionUnicode"
    return "Helvetica"


def export_pdf(report: str, evidence: Iterable[dict[str, Any]], filename: str) -> ReportExport:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Image as PdfImage
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("Thiếu reportlab; hãy cài lại requirements.txt.") from exc

    output = io.BytesIO()
    font = _register_pdf_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("MedVisionBody", parent=styles["BodyText"], fontName=font, leading=15)
    heading = ParagraphStyle("MedVisionHeading", parent=styles["Heading2"], fontName=font, spaceBefore=10)
    caption = ParagraphStyle("MedVisionCaption", parent=body, alignment=TA_CENTER, fontSize=8)
    doc = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    story: list[Any] = []
    report_body, review_gate = _split_review_gate(report)
    for raw_line in report_body.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 0.15 * cm))
            continue
        text = escape(line.lstrip("#- ").replace("**", "").replace("`", ""))
        story.append(Paragraph(text, heading if line.startswith("#") else body))
    story.extend(
        [
            PageBreak(),
            Paragraph("Phụ lục hình ảnh", heading),
            Paragraph(
                "Grad-CAM và pseudo bbox chỉ thể hiện vùng mô hình chú ý; không phải "
                "segmentation, annotation hoặc xác nhận tổn thương của bác sĩ.", body,
            ),
        ]
    )
    for item in evidence:
        story.append(Paragraph(escape(item["finding"]), heading))
        cells = []
        for kind, png in item["images"].items():
            image = PdfImage(io.BytesIO(png), width=8 * cm, height=8 * cm, kind="proportional")
            cells.append([image, Paragraph(escape(IMAGE_LABELS[kind]), caption)])
        rows = []
        for index in range(0, len(cells), 2):
            pair = cells[index:index + 2]
            while len(pair) < 2:
                pair.append(["", ""])
            rows.append([pair[0][0], pair[1][0]])
            rows.append([pair[0][1], pair[1][1]])
        table = Table(rows, colWidths=[8.5 * cm, 8.5 * cm])
        table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOX", (0, 0), (-1, -1), 0.25, colors.lightgrey)]))
        story.extend([table, Spacer(1, 0.35 * cm)])
    if review_gate:
        story.append(Paragraph("Trạng thái duyệt", heading))
        for line in review_gate.splitlines():
            story.append(Paragraph(escape(line), body))
    doc.build(story)
    return ReportExport(output.getvalue(), f"{filename}.pdf", "application/pdf")


def export_report(
    report: str,
    evidence: Iterable[dict[str, Any]],
    export_format: str,
    filename: str = "medvision_report",
) -> ReportExport:
    normalized = export_format.strip().upper()
    if normalized == "MARKDOWN (.MD)":
        return export_markdown(report, evidence, filename)
    if normalized == "PDF (.PDF)":
        return export_pdf(report, evidence, filename)
    if normalized == "WORD (.DOCX)":
        return export_docx(report, evidence, filename)
    raise ValueError(f"Định dạng xuất không được hỗ trợ: {export_format}")
